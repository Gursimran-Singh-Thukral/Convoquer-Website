"""Normalize the supplied Convoquer Teams workbook for participant CSV import."""
from __future__ import annotations

import csv
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports" / "Hard Coded Data" / "Convoquer Teams.xlsx"
OUTPUT = ROOT / "reports" / "Hard Coded Data" / "Convoquer Teams - Participant Import.csv"

INSTITUTES = {
    "Amity University Punjab": "Amity University Punjab",
    "MIET": "Model Institute of Engineering and Technology (MIET)",
    "CU": "Central University of Jammu",
    "ASCOMS": "Acharya Shiri Chander College of Medical Sciences and Hospital (ASCOMS)",
    "IIM Jammu": "IIM Jammu",
    "GMC": "GMC Jammu",
    "IIM Amritsar": "IIM Amritsar",
    "BGSBU": "Baba Ghulam Shah Badshah University, Rajouri",
    "GCET": "Government College of Engineering and Technology",
    "IIT Jammu": "IIT Jammu",
    "LPU": "Lovely Professional University",
}


def clean(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\ufffd", " ")).strip()


def key(value: str) -> str:
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def sport_of(value: str):
    raw = clean(value)
    k = key(raw.replace("Tournament Player Details", ""))
    gender = ""
    if not k:
        return None
    if "athletics" in k:
        gender = "FEMALE" if "w" in k.replace("athletics", "") else "MALE"
        return "Athletics", "", gender
    mappings = [
        ("badminton", "Badminton"), ("basketball", "Basketball"),
        ("tabletennis", "Table Tennis"), ("volleyball", "Volleyball"),
        ("chess", "Chess"), ("cricket", "Cricket"),
        ("football", "Football"), ("squash", "Squash"),
        ("weightlifting", "Weight Lifting"),
    ]
    for token, base in mappings:
        if token in k:
            suffix = k.replace(token, "")
            if base in {"Badminton", "Basketball", "Table Tennis", "Volleyball", "Chess"}:
                female = any(mark in suffix for mark in ("women", "woman", "girls", "girl")) or suffix == "w"
                gender = "FEMALE" if female else "MALE"
                return f"{base} ({'Women' if female else 'Men'})", "", gender
            return base, "", gender
    games = {"bgmi": "BGMI", "freefire": "Free Fire", "valorant": "Valorant"}
    if k in games:
        return "E-Sports", games[k], gender
    return None


def synthetic_roll(institute: str, name: str) -> str:
    return "TEMP-" + re.sub(r"[^A-Z0-9]+", "-", f"{institute}-{name}".upper()).strip("-")[:70]


def normalized_phone(value: str) -> str:
    digits = re.sub(r"\D", "", value)
    if len(digits) > 10 and digits.startswith("91"):
        digits = digits[-10:]
    return digits if 7 <= len(digits) <= 15 else ""


def add_row(rows, seen, *, name, college, roll, sport="", team="", gender="", phone="", role="PLAYER"):
    name = clean(name).replace(" - CAPTAIN", "").strip()
    if not name or key(name) in {"name", "playername"}:
        return
    roll = clean(roll) or synthetic_roll(college, name)
    identity = (key(college), key(roll), key(sport), key(team))
    if identity in seen:
        return
    seen.add(identity)
    rows.append({
        "name": name, "college": college, "rollNumber": roll,
        "sport": sport, "team": team, "gender": gender,
        "contactNumber": normalized_phone(phone), "role": role, "category": "ATHLETE",
    })


def parse_sheet(sheet, college: str, rows, seen):
    # ASCOMS has no sport/team markers. Preserve its people without inventing memberships.
    if sheet.title == "ASCOMS":
        for values in sheet.iter_rows(min_row=2, values_only=True):
            vals = [clean(v) for v in values]
            if len(vals) >= 3 and vals[1]:
                add_row(rows, seen, name=vals[1], college=college, roll=vals[2])
        return

    current = None
    captain = ""
    squad = ""
    header = {}
    for values in sheet.iter_rows(values_only=True):
        vals = [clean(v) for v in values]
        joined = " ".join(v for v in vals if v)

        detected = None
        # Prefer short standalone section cells; this avoids treating player data as headings.
        for value in vals:
            candidate = sport_of(value)
            if candidate and (len(value) < 80 or "Tournament" in value):
                detected = candidate
                break
        if detected:
            current = detected
            captain = ""
            squad = detected[1]
            header = {}

        lower = [v.lower() for v in vals]
        if any("captain name" in v for v in lower):
            candidates = [v for v in vals[1:] if v and not v.isdigit()]
            captain = candidates[0] if candidates else ""
            continue

        # GCET uses explicit Team 1/Team 2 lines inside each e-sport game.
        team_cells = [v for v in vals if re.fullmatch(r"Team\s+\d+", v, re.I)]
        if current and current[0] == "E-Sports" and team_cells:
            squad = f"{current[1]} - {team_cells[0]}"

        # Locate columns from each block's own header row.
        if any(key(v) in {"name", "playername"} for v in vals):
            for index, value in enumerate(vals):
                k = key(value)
                if k in {"name", "playername"}: header["name"] = index
                elif "studentid" in k or "registrationno" in k or "rollno" in k: header["roll"] = index
                elif "enrollmentno" in k: header["enrollment"] = index
                elif k in {"phno", "phoneno", "contactnumber"}: header["phone"] = index
            continue

        if not current:
            continue

        player_marker = next((i for i, v in enumerate(vals) if re.match(r"^Player\s*\d+", v, re.I)), None)
        is_athletics_row = current[0] == "Athletics" and len(vals) > 2 and vals[2]
        is_weight_row = current[0] == "Weight Lifting" and len(vals) > 2 and vals[2]
        if player_marker is None and not is_athletics_row and not is_weight_row:
            continue

        name_index = header.get("name", 2 if is_athletics_row or is_weight_row else player_marker + 1)
        name = vals[name_index] if name_index < len(vals) else ""
        if not name:
            continue

        # Some weightlifting cells contain comma-separated names without IDs.
        names = [clean(n) for n in name.split(",")] if is_weight_row and "," in name else [name]
        roll_index = header.get("enrollment", header.get("roll", name_index + 1))
        roll = vals[roll_index] if roll_index < len(vals) else ""
        # CU's first identifier is often email; its Enrollment No is the stable roll number.
        if "@" in roll and len(vals) > 5 and vals[5]:
            roll = vals[5]
        phone_index = header.get("phone", -1)
        phone = vals[phone_index] if 0 <= phone_index < len(vals) else ""
        for person in names:
            add_row(
                rows, seen, name=person, college=college, roll=roll,
                sport=current[0], team=squad if current[0] == "E-Sports" else "",
                gender=current[2], phone=phone,
                role="CAPTAIN" if captain and key(person) == key(captain) else "PLAYER",
            )


def main() -> int:
    if not SOURCE.exists():
        print(f"Workbook not found: {SOURCE}", file=sys.stderr)
        return 1
    workbook = openpyxl.load_workbook(SOURCE, data_only=True)
    rows, seen = [], set()
    for sheet_name, college in INSTITUTES.items():
        if sheet_name in workbook.sheetnames:
            parse_sheet(workbook[sheet_name], college, rows, seen)
    rows.sort(key=lambda row: (row["college"].lower(), row["sport"].lower(), row["team"].lower(), row["name"].lower()))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    fields = ["name", "college", "rollNumber", "sport", "team", "gender", "contactNumber", "role", "category"]
    with OUTPUT.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    print(f"Wrote {len(rows)} membership rows to {OUTPUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
